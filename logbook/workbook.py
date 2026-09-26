"""Writing the spreadsheet: plain, ruled, no colour, one sheet per financial year
when the driver asks for that.

Anything a person typed is written as plain text, never as a formula, so a
registration entered as =HYPERLINK(...) cannot run when the file is opened. Only
the cells this module builds itself are live formulas.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .dates import financial_year_bounds, financial_year_of, next_financial_year
from .errors import fail
from .text import sanitize_sheet_name

FONT_NAME = "Arial"
BASE_SIZE = 10
LAST_COL = 9                      # A..I

THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
UNDERLINE = Border(bottom=THIN)

DATE_FMT = "DD/MM/YYYY"
KM_FMT = "#,##0"
MONEY_FMT = "$#,##0.00"
PCT_FMT = "0.0%"

HEADERS = ["No.", "Journey start date", "Journey end date", "Purpose",
           "Reason for journey", "Destination", "Odometer start", "Odometer end",
           "Kilometres"]
COLUMN_WIDTHS = [6, 18, 18, 12, 44, 30, 15, 15, 12]


def fixed_timestamps(raw: bytes, stamp) -> bytes:
    """Rewrite the workbook's zip entries with one fixed date.

    A spreadsheet is a zip, and each entry normally carries the moment it was
    written. Pinning them keeps the file free of a wall-clock trace and means the
    same answers always produce the same bytes.
    """
    when = (max(1980, stamp.year), stamp.month, stamp.day, 0, 0, 0)
    fixed = f"{stamp:%Y-%m-%d}T00:00:00Z"
    source = zipfile.ZipFile(io.BytesIO(raw))
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</)",
                                 lambda m: m.group(1) + fixed.encode() + m.group(2),
                                 content)
            info = zipfile.ZipInfo(item.filename, date_time=when)
            info.compress_type = item.compress_type
            info.external_attr = item.external_attr
            info.create_system = 0
            target.writestr(info, content)
    return buffer.getvalue()


def put(ws, row, col, value=None, *, bold=False, size=BASE_SIZE, fmt=None,
        align=None, wrap=False, border=None, italic=False, formula=False):
    """Write a cell. Anything the user supplied is forced to plain text."""
    cell = ws.cell(row=row, column=col, value=value)
    if not formula and isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    cell.font = Font(name=FONT_NAME, size=size, bold=bold, italic=italic)
    cell.alignment = Alignment(horizontal=align or "left", vertical="center",
                               wrap_text=wrap)
    if fmt:
        cell.number_format = fmt
    if border:
        cell.border = border
    return cell


def meta_line(ws, row, label, value, *, fmt=None, bold=False, formula=False):
    put(ws, row, 1, label, bold=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    put(ws, row, 4, value, fmt=fmt, bold=bold, formula=formula)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=LAST_COL)
    return row + 1


def heading(ws, row, text):
    put(ws, row, 1, text, bold=True, size=BASE_SIZE + 1, border=UNDERLINE)
    for col in range(2, LAST_COL + 1):
        put(ws, row, col, None, border=UNDERLINE)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST_COL)
    return row + 1


def fill_sheet(ws, cfg, rows, start: date, end: date,
               start_odo: int, end_odo: int, income: float):
    """One complete logbook on one worksheet."""
    row = 1
    put(ws, row, 1, "MOTOR VEHICLE LOGBOOK", bold=True, size=14)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST_COL)
    row += 1
    put(ws, row, 1, "Logbook method - record of business and private travel",
        size=9, italic=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=LAST_COL)
    row += 2

    row = heading(ws, row, "Vehicle and owner")
    if cfg.owner_name:
        row = meta_line(ws, row, "Driver / owner", cfg.owner_name)
    if cfg.abn:
        row = meta_line(ws, row, "ABN", cfg.abn)
    row = meta_line(ws, row, "Make", cfg.make)
    row = meta_line(ws, row, "Model", cfg.model)
    if cfg.year_of_manufacture:
        row = meta_line(ws, row, "Year of manufacture", cfg.year_of_manufacture)
    row = meta_line(ws, row, "Engine capacity", cfg.engine_capacity)
    row = meta_line(ws, row, "Registration number", cfg.registration)
    row += 1

    row = heading(ws, row, "Logbook period and odometer")
    row = meta_line(ws, row, "Period begins", start, fmt=DATE_FMT)
    period_end_row = row
    row = meta_line(ws, row, "Period ends", end, fmt=DATE_FMT)
    odo_start_row = row
    row = meta_line(ws, row, "Odometer reading at start of period", start_odo, fmt=KM_FMT)
    odo_end_row = row
    row = meta_line(ws, row, "Odometer reading at end of period", end_odo, fmt=KM_FMT)
    total_km_row = row
    row = meta_line(ws, row, "Total kilometres travelled in period",
                    f"=D{odo_end_row}-D{odo_start_row}", fmt=KM_FMT, bold=True,
                    formula=True)
    row += 1

    row = heading(ws, row, "Business use")
    row = meta_line(ws, row, "Income earned in the period", income, fmt=MONEY_FMT)
    business_row = row
    row = meta_line(ws, row, "Business kilometres", None, fmt=KM_FMT)
    private_row = row
    row = meta_line(ws, row, "Private kilometres", None, fmt=KM_FMT)
    pct_row = row
    row = meta_line(ws, row, "BUSINESS-USE PERCENTAGE", None, fmt=PCT_FMT, bold=True)
    row = meta_line(ws, row, "Logbook valid until (5 years)",
                    f"=DATE(YEAR(D{period_end_row})+5,MONTH(D{period_end_row}),"
                    f"DAY(D{period_end_row}))", fmt=DATE_FMT, formula=True)
    row += 1

    header_row = row
    for i, text in enumerate(HEADERS, start=1):
        put(ws, header_row, i, text, bold=True, align="center", wrap=True, border=BOX)
    ws.row_dimensions[header_row].height = 30

    first = header_row + 1
    for i, entry in enumerate(rows):
        at = first + i
        put(ws, at, 1, entry["n"], align="center", border=BOX)
        put(ws, at, 2, entry["date"], fmt=DATE_FMT, align="center", border=BOX)
        put(ws, at, 3, entry["date"], fmt=DATE_FMT, align="center", border=BOX)
        put(ws, at, 4, entry["purpose"], align="center", border=BOX)
        put(ws, at, 5, entry["reason"], border=BOX)
        put(ws, at, 6, entry.get("destination", ""), border=BOX)
        put(ws, at, 7, entry["odo_start"], fmt=KM_FMT, align="right", border=BOX)
        put(ws, at, 8, entry["odo_end"], fmt=KM_FMT, align="right", border=BOX)
        put(ws, at, 9, f"=H{at}-G{at}", fmt=KM_FMT, align="right", border=BOX,
            formula=True)
    last = first + len(rows) - 1

    total_row = last + 1
    put(ws, total_row, 1, "Total", bold=True, border=BOX)
    for col in range(2, LAST_COL):
        put(ws, total_row, col, None, border=BOX)
    put(ws, total_row, 6, "Total kilometres", bold=True, align="right", border=BOX)
    put(ws, total_row, 9, f"=SUM(I{first}:I{last})", bold=True, fmt=KM_FMT,
        align="right", border=BOX, formula=True)

    business_total_row = total_row + 1
    put(ws, business_total_row, 6, "Business kilometres", bold=True, align="right",
        border=BOX)
    put(ws, business_total_row, 9,
        f'=SUMIF($D${first}:$D${last},"Business",$I${first}:$I${last})',
        bold=True, fmt=KM_FMT, align="right", border=BOX, formula=True)
    private_total_row = total_row + 2
    put(ws, private_total_row, 6, "Private kilometres", bold=True, align="right",
        border=BOX)
    put(ws, private_total_row, 9,
        f'=SUMIF($D${first}:$D${last},"Private",$I${first}:$I${last})',
        bold=True, fmt=KM_FMT, align="right", border=BOX, formula=True)
    pct_total_row = total_row + 3
    put(ws, pct_total_row, 6, "Business-use percentage", bold=True, align="right",
        border=BOX)
    put(ws, pct_total_row, 9, f"=IFERROR(I{business_total_row}/I{total_row},0)",
        bold=True, fmt=PCT_FMT, align="right", border=BOX, formula=True)

    ws.cell(row=business_row, column=4).value = f"=I{business_total_row}"
    ws.cell(row=private_row, column=4).value = f"=I{private_total_row}"
    ws.cell(row=pct_row, column=4).value = f"=IFERROR(D{business_row}/D{total_km_row},0)"

    row = pct_total_row + 3
    row = heading(ws, row, "Odometer readings for each income year")
    for col, text in ((2, "Income year"), (3, "Odometer at start of year"),
                      (4, "Odometer at end of year"), (5, "Total km for the year"),
                      (6, "Business km at the logbook percentage")):
        put(ws, row, col, text, bold=True, align="center", wrap=True, border=BOX)
    ws.row_dimensions[row].height = 30
    row += 1
    label = financial_year_of(start)
    for i in range(5):
        year_start, year_end = financial_year_bounds(label)
        put(ws, row, 2, label, align="center", border=BOX)
        opening = start_odo if (i == 0 and start == year_start) else None
        closing = end_odo if (i == 0 and end == year_end) else None
        put(ws, row, 3, opening, fmt=KM_FMT, align="right", border=BOX)
        put(ws, row, 4, closing, fmt=KM_FMT, align="right", border=BOX)
        put(ws, row, 5, f'=IF(COUNT(C{row}:D{row})=2,D{row}-C{row},"")', fmt=KM_FMT,
            align="right", border=BOX, formula=True)
        put(ws, row, 6, f'=IF(ISNUMBER(E{row}),ROUND(E{row}*$D${pct_row},0),"")',
            fmt=KM_FMT, align="right", border=BOX, formula=True)
        label = next_financial_year(label)
        row += 1

    for i, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_workbook(cfg, sections: list, out):
    """Write one worksheet per section. `out` may be a path or an open binary file.

    A section is (sheet name, rows, start, end, opening odometer, closing odometer,
    income). One section gives the ordinary single-sheet logbook.
    """
    if not sections:
        fail("There is nothing to write - no journeys were produced.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    used = set()
    for name, rows, start, end, start_odo, end_odo, income in sections:
        title = sanitize_sheet_name(name)
        base, n = title, 2
        while title in used:                     # Excel will not take two alike
            suffix = f" {n}"
            title = base[:31 - len(suffix)] + suffix
            n += 1
        used.add(title)
        fill_sheet(workbook.create_sheet(title), cfg, rows, start, end,
                   start_odo, end_odo, income)

    stamp = datetime(sections[-1][3].year, sections[-1][3].month, sections[-1][3].day)
    workbook.properties.created = stamp
    workbook.properties.modified = stamp
    workbook.properties.creator = "Car logbook generator"
    workbook.properties.lastModifiedBy = "Car logbook generator"
    workbook.properties.title = "Motor vehicle logbook"
    workbook.calculation.fullCalcOnLoad = True

    raw = io.BytesIO()
    workbook.save(raw)
    data = fixed_timestamps(raw.getvalue(), stamp)
    if hasattr(out, "write"):
        out.write(data)
    else:
        with open(out, "wb") as handle:
            handle.write(data)
    return data


def workbook_bytes(cfg, sections: list) -> bytes:
    return write_workbook(cfg, sections, io.BytesIO())
