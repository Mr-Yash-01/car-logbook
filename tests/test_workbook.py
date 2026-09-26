#!/usr/bin/env python3
"""The spreadsheet that comes out: layout, formulas, the financial-year split,
and everything a person typed staying plain text."""

from __future__ import annotations

import sys
import tempfile
import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from harness import check, group, refuses, report          # noqa: E402

from openpyxl import load_workbook                         # noqa: E402

import logbook as L                                        # noqa: E402
from logbook.config import Config, normalise_config        # noqa: E402

TMP = Path(tempfile.mkdtemp())
SETTINGS = dict(start_date="2025-07-01", end_date="2025-09-22",
                start_odometer=45230, end_odometer=55230, work_type="both",
                work_days=[4, 5, 6], business_share_on_work_days=0.85,
                prorate_partial_months=False, owner_name="Test Driver",
                abn="51 824 753 556", make="Toyota", model="Corolla",
                year_of_manufacture="2021", engine_capacity="1.8L",
                registration="ABC123",
                monthly_income={"2025-07": 5000, "2025-08": 5000, "2025-09": 3667})


def sheet_rows(ws):
    return [r for r in ws.iter_rows(values_only=True)
            if isinstance(r[0], int) and isinstance(r[6], int)]


def sheet_text(ws):
    return " ".join(str(c.value) for row in ws.iter_rows()
                    for c in row if c.value is not None)


# ===========================================================================
group("A. One sheet, the ordinary case")
# ===========================================================================

result = L.generate(SETTINGS)
book = load_workbook(BytesIO(result["bytes"]))
ws = book.active
text = sheet_text(ws)

check("one sheet", len(book.sheetnames) == 1, book.sheetnames)
check("nine columns", ws.max_column == 9, ws.max_column)
check("it really is a spreadsheet", result["bytes"][:2] == b"PK")
check("filename carries the plate and dates",
      result["filename"] == "logbook-ABC123-2025-07-01-to-2025-09-22.xlsx",
      result["filename"])
for field in ["Test Driver", "51 824 753 556", "2021", "Toyota", "Corolla",
              "1.8L", "ABC123"]:
    check(f"the sheet shows {field!r}", field in text)
check("no colour anywhere",
      not any(c.fill and c.fill.fgColor and c.fill.fgColor.rgb not in (None, "00000000")
              for row in ws.iter_rows() for c in row))
check("one font throughout",
      {c.font.name for row in ws.iter_rows() for c in row if c.value is not None}
      == {"Arial"})
check("no weekday column", not any(d in text for d in L.DAY_NAMES))
check("the income-year block is there", "Income year" in text)
check("no rate per kilometre is shown",
      "per business kilometre" not in text and "Dollars earned" not in text)
check("no weeks-covered row", "Continuous weeks" not in text)

header_row = next(r[0].row for r in ws.iter_rows(max_col=1) if r[0].value == "No.")
headers = [c.value for c in ws[header_row]]
check("the columns are in order",
      headers == ["No.", "Journey start date", "Journey end date", "Purpose",
                  "Reason for journey", "Destination", "Odometer start",
                  "Odometer end", "Kilometres"], headers)

rows = sheet_rows(ws)
check("journeys were written", len(rows) > 40, len(rows))
check("it opens on the reading given", rows[0][6] == 45230)
check("it closes on the reading given", rows[-1][7] == 55230)
check("the odometer never jumps",
      all(rows[i][7] == rows[i + 1][6] for i in range(len(rows) - 1)))
check("both dates on every row", all(r[1] == r[2] for r in rows))
check("purposes are only business or private", {r[3] for r in rows} <= {"Business", "Private"})
check("every work journey names an area",
      all(str(r[5] or "").strip() for r in rows if r[3] == "Business"))
check("private journeys name none",
      all(not str(r[5] or "").strip() for r in rows if r[3] == "Private"))
check("kilometres are formulas",
      all(isinstance(r[8], str) and r[8].startswith("=") for r in rows))
check("business never lands on a day off",
      all(r[1].weekday() in (4, 5, 6) for r in rows if r[3] == "Business"))
check("more than one wording is used",
      len({r[4] for r in rows if r[3] == "Business"}) > 2)
check("more than one area is used", len({r[5] for r in rows if r[3] == "Business"}) > 2)


# ===========================================================================
group("B. The sums")
# ===========================================================================

totals_row = next(r[0].row for r in ws.iter_rows(max_col=1) if r[0].value == "Total")
first = header_row + 1
last = totals_row - 1
check("the total spans every row",
      ws.cell(row=totals_row, column=9).value == f"=SUM(I{first}:I{last})",
      ws.cell(row=totals_row, column=9).value)
check("business uses SUMIF", "SUMIF" in ws.cell(row=totals_row + 1, column=9).value)
check("the percentage divides business by total",
      ws.cell(row=totals_row + 3, column=9).value
      == f"=IFERROR(I{totals_row + 1}/I{totals_row},0)")
check("shown as a percentage",
      ws.cell(row=totals_row + 3, column=9).number_format == "0.0%")
check("odometer cells carry thousands separators",
      ws.cell(row=first, column=7).number_format == "#,##0")
check("dates are day/month/year", ws.cell(row=first, column=2).number_format == "DD/MM/YYYY")
check("the header repeats when printed", f"${header_row}:${header_row}"
      in str(ws.print_title_rows) or str(ws.print_title_rows) == f"{header_row}:{header_row}")


def label_row(sheet, text):
    for row in sheet.iter_rows(max_col=1):
        if row[0].value and str(row[0].value).startswith(text):
            return row[0].row
    return None


begins = label_row(ws, "Period begins")
ends = label_row(ws, "Period ends")
odo_start = label_row(ws, "Odometer reading at start")
odo_end = label_row(ws, "Odometer reading at end")
total_km = label_row(ws, "Total kilometres travelled")
valid = label_row(ws, "Logbook valid until")
check("valid-until counts from the period end",
      f"D{ends}" in ws.cell(row=valid, column=4).value
      and f"D{begins}" not in ws.cell(row=valid, column=4).value)
check("total kilometres subtracts the two readings",
      ws.cell(row=total_km, column=4).value == f"=D{odo_end}-D{odo_start}")
check("the rows are in order", begins < ends < odo_start < odo_end < total_km)


# ===========================================================================
group("C. A sheet for each financial year")
# ===========================================================================

across = dict(SETTINGS, start_date="2025-05-01", end_date="2025-08-31",
              start_odometer=20000, end_odometer=32000,
              monthly_income={f"2025-{m:02d}": 4000 for m in (5, 6, 7, 8)})
check("the period does cross 30 June",
      L.spans_financial_years(date(2025, 5, 1), date(2025, 8, 31)))

combined = L.generate(dict(across, sheet_mode="combined"))
check("combined keeps one sheet", combined["sheets"] == ["Logbook"], combined["sheets"])

split = L.generate(dict(across, sheet_mode="per-year"))
check("per-year gives a sheet each",
      split["sheets"] == ["2024-25 logbook", "2025-26 logbook"], split["sheets"])
book2 = load_workbook(BytesIO(split["bytes"]))
first_sheet, second_sheet = book2.worksheets
rows_a, rows_b = sheet_rows(first_sheet), sheet_rows(second_sheet)
check("both sheets have journeys", rows_a and rows_b)
check("the first opens on the starting reading", rows_a[0][6] == 20000)
check("the second closes on the closing reading", rows_b[-1][7] == 32000)
check("the second picks up where the first left off", rows_a[-1][7] == rows_b[0][6])
check("each sheet chains cleanly",
      all(rows_a[i][7] == rows_a[i + 1][6] for i in range(len(rows_a) - 1))
      and all(rows_b[i][7] == rows_b[i + 1][6] for i in range(len(rows_b) - 1)))
check("together they hold every journey", len(rows_a) + len(rows_b) ==
      len(sheet_rows(load_workbook(BytesIO(combined["bytes"])).active)))
check("each sheet numbers its rows from one",
      rows_a[0][0] == 1 and rows_b[0][0] == 1)
check("the first sheet ends on 30 June",
      max(r[1] for r in rows_a).date() <= date(2025, 6, 30))
check("the second starts on 1 July",
      min(r[1] for r in rows_b).date() >= date(2025, 7, 1))
check("asking for per-year on a single year changes nothing",
      L.generate(dict(SETTINGS, sheet_mode="per-year"))["sheets"] == ["Logbook"])


# ===========================================================================
group("D. Nothing a person typed can run")
# ===========================================================================

cfg = Config(**{k: v for k, v in SETTINGS.items() if k in Config.__dataclass_fields__})
cfg.rideshare_reasons = ['=HYPERLINK("http://evil.example","click")', "Airport run"]
cfg.rideshare_destinations = ["=1+1", "Brisbane CBD"]
cfg.delivery_reasons = ["=cmd|'/c calc'!A0", "Food run"]
normalise_config(cfg)
danger = L.generate(cfg)
ws3 = load_workbook(BytesIO(danger["bytes"])).active
live = [c.value for row in ws3.iter_rows() for c in row
        if c.data_type == "f" and ("HYPERLINK" in str(c.value) or "cmd|" in str(c.value))]
check("a formula typed as wording is stored as text", live == [], live[:1])
check("our own totals are still live formulas",
      sum(1 for row in ws3.iter_rows() for c in row if c.data_type == "f") > 50)
for bad in ("=HYPERLINK(1)", "<script>", "x" * 60):
    check(f"a make of {bad[:14]!r} is refused outright",
          refuses(L.generate, dict(SETTINGS, make=bad))[0])
check("a plate of ABCDEFGH is refused",
      refuses(L.generate, dict(SETTINGS, registration="ABCDEFGH"))[0])


# ===========================================================================
group("E. The same answers give the same file")
# ===========================================================================

import time                                              # noqa: E402
one = L.generate(SETTINGS)["bytes"]
time.sleep(1.1)
two = L.generate(SETTINGS)["bytes"]
check("identical across a second boundary", one == two)
check("a different shuffle number changes it",
      L.generate(dict(SETTINGS, seed=7))["bytes"] != one)
props = load_workbook(BytesIO(one)).properties
check("no wall-clock timestamp inside",
      props.created.date() == date(2025, 9, 22) and props.modified.date() == date(2025, 9, 22),
      (props.created, props.modified))
check("nor on the zip entries",
      {i.date_time[3:] for i in zipfile.ZipFile(BytesIO(one)).infolist()} == {(0, 0, 0)})
check("it recalculates as soon as it opens",
      load_workbook(BytesIO(one)).calculation.fullCalcOnLoad is True)

path = TMP / "written.xlsx"
L.write_workbook(Config(), [("Logbook", L.build_journeys(
    Config(work_days=[4, 5, 6], rideshare_reasons=["A", "B"],
           rideshare_destinations=["X", "Y"], private_reasons=["P"]),
    date(2025, 7, 1), date(2025, 9, 22), 0, 10000,
    {"2025-07": 2500, "2025-08": 2500, "2025-09": 1800})[0],
    date(2025, 7, 1), date(2025, 9, 22), 0, 10000, 1.0)], str(path))
check("writing straight to a path works", path.exists() and path.read_bytes()[:2] == b"PK")
refused, message = refuses(L.write_workbook, Config(), [], BytesIO())
check("an empty workbook is refused in plain words",
      refused and "nothing to write" in message, message)

report("Workbook")
