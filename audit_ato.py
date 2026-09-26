#!/usr/bin/env python3
"""
ATO compliance audit for a generated logbook
============================================

Checks a workbook against the record-keeping requirements on the ATO's
"Logbook method" page (ato.gov.au/logbook), item by item.

    python audit_ato.py sample_logbook.xlsx

PASS   the requirement is met by the file
GAP    a required item is missing
REVIEW the item is present but only you can judge whether it is good enough

This checks the shape of the record, not the truth of it. A logbook must
describe journeys actually made; nothing here can verify that.
"""

from __future__ import annotations

import sys
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook

RESULTS = []
SPAN = []          # every sheet's period, so the 12-week rule is judged once


def report(status, requirement, detail=""):
    RESULTS.append((status, requirement, detail))


def purpose_col_guess(cols):
    return cols.index("Purpose") if "Purpose" in cols else None


def label_value(ws, prefix):
    """Find a metadata row by its label and return its value."""
    for row in ws.iter_rows(max_col=4):
        if row[0].value and str(row[0].value).strip().lower().startswith(prefix.lower()):
            return row[3].value
    return None


def sheet_text(ws):
    return " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    return v if isinstance(v, date) else None


def number_or_none(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def audit(path, ws=None, check_weeks=True):
    # Values where the file has them; a fresh download has formulas but no stored
    # results until a spreadsheet app opens it, so anything missing is worked out
    # from the odometer columns, which are always plain numbers.
    if ws is None:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
    text = sheet_text(ws)

    hdr = None
    for row in ws.iter_rows(max_col=1):
        if row[0].value == "No.":
            hdr = row[0].row
            break
    if hdr is None:
        report("GAP", "A table of journeys", "no journey table found in the sheet")
        return
    cols = [c.value for c in ws[hdr]]
    odo_idx = cols.index("Odometer start") if "Odometer start" in cols else None
    rows = [r for r in ws.iter_rows(min_row=hdr + 1, values_only=True)
            if isinstance(r[0], int)
            and (odo_idx is None or isinstance(r[odo_idx], (int, float)))]

    # ---- the logbook must contain -----------------------------------------
    begins, ends = as_date(label_value(ws, "Period begins")), as_date(label_value(ws, "Period ends"))
    if begins and ends and ends > begins:
        report("PASS", "When the logbook period begins and ends",
               f"{begins:%d/%m/%Y} to {ends:%d/%m/%Y}")
    else:
        report("GAP", "When the logbook period begins and ends", "missing or out of order")

    odo_s = label_value(ws, "Odometer reading at start")
    odo_e = label_value(ws, "Odometer reading at end")
    if isinstance(odo_s, (int, float)) and isinstance(odo_e, (int, float)) and odo_e > odo_s:
        report("PASS", "Odometer readings at the start and end of the period",
               f"{odo_s:,.0f} to {odo_e:,.0f}")
    else:
        report("GAP", "Odometer readings at the start and end of the period", "missing")

    total = number_or_none(label_value(ws, "Total kilometres travelled"))
    if total is None and isinstance(odo_s, (int, float)) and isinstance(odo_e, (int, float)):
        total = odo_e - odo_s
    if isinstance(total, (int, float)) and total == (odo_e or 0) - (odo_s or 0):
        report("PASS", "Total kilometres travelled during the period",
               f"{total:,.0f} km, matching the odometer readings")
    else:
        report("GAP", "Total kilometres travelled during the period",
               f"shows {total}, odometer difference is {(odo_e or 0) - (odo_s or 0)}")

    km_col = cols.index("Kilometres") if "Kilometres" in cols else None
    end_idx = cols.index("Odometer end") if "Odometer end" in cols else None

    def journey_km(row):
        """Distance for a row: the stored figure, else the odometer difference."""
        if km_col is not None:
            direct = number_or_none(row[km_col])
            if direct is not None:
                return direct
        if odo_idx is not None and end_idx is not None:
            a, b = number_or_none(row[odo_idx]), number_or_none(row[end_idx])
            if a is not None and b is not None:
                return b - a
        return None

    distances = [journey_km(r) for r in rows]
    if km_col is not None and rows and all(d is not None and d > 0 for d in distances):
        report("PASS", "Kilometres travelled on each journey",
               f"{len(rows)} journeys, each with its own distance")
    else:
        report("GAP", "Kilometres travelled on each journey", "column missing or incomplete")

    iy_hdr = None
    for row in ws.iter_rows(max_col=2):
        if row[1].value == "Income year":
            iy_hdr = row[0].row
            break
    if iy_hdr:
        years, filled, cells = 0, 0, 0
        for row in ws.iter_rows(min_row=iy_hdr + 1, max_col=4, values_only=True):
            if not row[1]:
                break
            years += 1
            cells += 2
            filled += sum(1 for v in (row[2], row[3]) if isinstance(v, (int, float)))
        report("PASS", "Odometer readings at the start and end of each later income year",
               f"{years} income years listed; {filled} of {cells} readings recorded so far "
               f"- add each one as the year turns")
    elif any(k in text for k in ("1 July", "30 June", "income year")):
        report("PASS", "Odometer readings at the start and end of each later income year")
    else:
        report("GAP", "Odometer readings at the start and end of each later income year",
               "not in this sheet - required for every income year the logbook is relied on")

    pct = number_or_none(label_value(ws, "BUSINESS-USE PERCENTAGE"))
    if pct is None and total and purpose_col_guess(cols) is not None:
        pcol = purpose_col_guess(cols)
        business_total = sum(d for r, d in zip(rows, distances)
                             if d and r[pcol] == "Business")
        pct = business_total / total if total else None
    if isinstance(pct, (int, float)) and 0 < pct <= 1:
        report("PASS", "Business-use percentage for the logbook period", f"{pct * 100:.1f}%")
    else:
        report("GAP", "Business-use percentage for the logbook period", f"shows {pct}")

    details = {k: label_value(ws, k) for k in
               ("Make", "Model", "Engine capacity", "Registration number")}
    missing = [k for k, v in details.items() if not v]
    if not missing:
        report("PASS", "Make, model, engine capacity and registration number",
               ", ".join(str(v) for v in details.values()))
    else:
        report("GAP", "Make, model, engine capacity and registration number",
               f"missing: {', '.join(missing)}")

    # ---- for each journey --------------------------------------------------
    reason_col = cols.index("Reason for journey") if "Reason for journey" in cols else None
    purpose_col = cols.index("Purpose") if "Purpose" in cols else None
    if reason_col is not None and all(str(r[reason_col] or "").strip() for r in rows):
        reasons = {str(r[reason_col]).strip() for r in rows}
        thin = {r for r in reasons if r.lower() in
                ("business", "miscellaneous business", "private", "work", "personal")}
        if thin:
            report("GAP", "Reason for each journey",
                   f"entries such as {sorted(thin)[0]!r} are too bare to stand on their own")
        else:
            biz = {str(r[reason_col]).strip() for r in rows
                   if purpose_col is not None and r[purpose_col] == "Business"}
            report("REVIEW" if len(biz) < 2 else "PASS", "Reason for each journey",
                   f"{len(reasons)} distinct wordings ({len(biz)} for business)"
                   + ("; the same line repeats on every business journey"
                      if len(biz) < 2 else ""))
    else:
        report("GAP", "Reason for each journey", "column missing or blank on some rows")

    if "Journey start date" in cols and "Journey end date" in cols:
        sd, ed = cols.index("Journey start date"), cols.index("Journey end date")
        if all(r[sd] and r[ed] for r in rows):
            report("PASS", "Start and end date of each journey",
                   "both dates on every row (same-day journeys combined, which is allowed)")
        else:
            report("GAP", "Start and end date of each journey", "blank on some rows")
    else:
        report("GAP", "Start and end date of each journey", "columns missing")

    if "Odometer start" in cols and "Odometer end" in cols:
        os_, oe_ = cols.index("Odometer start"), cols.index("Odometer end")
        if all(isinstance(r[os_], (int, float)) and isinstance(r[oe_], (int, float))
               for r in rows):
            report("PASS", "Odometer readings at the start and end of each journey")
        else:
            report("GAP", "Odometer readings at the start and end of each journey",
                   "blank on some rows")
    else:
        report("GAP", "Odometer readings at the start and end of each journey",
               "columns missing")

    # ---- timeframe ---------------------------------------------------------
    # Split across financial years, the 12-week rule applies to the whole period,
    # not to one year's sheet, so it is checked once at the end instead.
    if begins and ends and check_weeks:
        days = (ends - begins).days + 1
        if days >= 84:
            report("PASS", "At least 12 continuous weeks",
                   f"{days} days ({days / 7:.1f} weeks)")
        else:
            report("GAP", "At least 12 continuous weeks",
                   f"only {days} days ({days / 7:.1f} weeks)")
    if begins and ends:
        SPAN.append((begins, ends))

    # ---- integrity of the record ------------------------------------------
    if rows and "Odometer start" in cols:
        os_, oe_ = cols.index("Odometer start"), cols.index("Odometer end")
        broken = [i for i in range(len(rows) - 1) if rows[i][oe_] != rows[i + 1][os_]]
        opens = rows[0][os_] == odo_s
        closes = rows[-1][oe_] == odo_e
        if not broken and opens and closes:
            report("PASS", "Unbroken odometer trail",
                   "every journey starts where the previous one ended, opening to closing "
                   "reading, so no travel is unaccounted for")
        else:
            report("GAP", "Unbroken odometer trail",
                   f"{len(broken)} break(s); opens correctly: {opens}; closes correctly: {closes}")

    if rows and begins and ends and "Journey start date" in cols:
        sd = cols.index("Journey start date")
        dates = [as_date(r[sd]) for r in rows]
        blanks = sum(1 for d in dates if d is None)
        present = [d for d in dates if d is not None]
        if blanks or not present:
            report("GAP", "Journeys sit inside the period",
                   f"{blanks} row(s) carry no date")
        else:
            inside = all(begins <= d <= ends for d in present)
            ordered = present == sorted(present)
            if inside and ordered:
                report("PASS", "Journeys sit inside the period",
                       f"{min(present):%d/%m/%Y} to {max(present):%d/%m/%Y}, in date order")
            else:
                report("GAP", "Journeys sit inside the period",
                       ("some journeys fall outside the period" if not inside
                        else "journeys are out of date order"))

    if rows and purpose_col is not None:
        bus = sum(d for r, d in zip(rows, distances) if d and r[purpose_col] == "Business")
        pri = sum(d for r, d in zip(rows, distances) if d and r[purpose_col] == "Private")
        if isinstance(total, (int, float)) and bus + pri == total:
            report("PASS", "Private travel recorded as well as business",
                   f"{bus:,.0f} business + {pri:,.0f} private = {total:,.0f} km")
        else:
            report("GAP", "Private travel recorded as well as business",
                   f"{bus:,.0f} + {pri:,.0f} does not reach {total}")

    # ---- things the sheet cannot show -------------------------------------
    # Destination is not on the ATO's required list for a business logbook, so this
    # is reported as supporting detail rather than a gap.
    if "Destination" in cols:
        dcol = cols.index("Destination")
        filled = sum(1 for r in rows if str(r[dcol] or "").strip())
        work = sum(1 for r in rows if purpose_col is not None and r[purpose_col] == "Business")
        if work and filled >= work:
            report("PASS", "Destination of each work journey (supporting detail)",
                   f"all {work} work journeys name an area")
        elif filled:
            report("REVIEW", "Destination of each work journey (supporting detail)",
                   f"only {filled} of {work} work journeys name an area")
        else:
            report("REVIEW", "Destination of each work journey (supporting detail)",
                   "column is there but empty")


def main():
    if len(sys.argv) < 2:
        sys.exit("Usage: python audit_ato.py <logbook.xlsx>")
    path = Path(sys.argv[1])
    if not path.exists():
        sys.exit(f"No file at {path}")
    try:
        book = load_workbook(path, data_only=True)
        sheets = book.worksheets
        many = len(sheets) > 1
        for sheet in sheets:
            if many:
                report("SHEET", sheet.title, "")
            audit(path, sheet, check_weeks=not many)
        if many and SPAN:
            first = min(s for s, _e in SPAN)
            last = max(e for _s, e in SPAN)
            days = (last - first).days + 1
            report("PASS" if days >= 84 else "GAP",
                   "At least 12 continuous weeks, across every sheet",
                   f"{first:%d/%m/%Y} to {last:%d/%m/%Y} - {days} days "
                   f"({days / 7:.1f} weeks)")
    except Exception as exc:                                    # noqa: BLE001
        report("GAP", "The file could be read end to end",
               f"{type(exc).__name__}: {exc}")

    width = max(len(r[1]) for r in RESULTS) + 2
    print(f"\nATO logbook audit: {path.name}")
    print("=" * (width + 46))
    for status, req, detail in RESULTS:
        print(f"  {status:<7} {req:<{width}} {detail}")
    print("=" * (width + 46))
    counts = {s: sum(1 for r in RESULTS if r[0] == s) for s in ("PASS", "GAP", "REVIEW")}
    print(f"  {counts['PASS']} pass, {counts['GAP']} gap, {counts['REVIEW']} to review")
    print("\n  Source: ato.gov.au/logbook. A logbook must record journeys actually made;")
    print("  this audit checks the record's shape, not whether the travel happened.")
    return 1 if counts["GAP"] else 0


if __name__ == "__main__":
    sys.exit(main())
